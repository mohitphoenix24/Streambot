"""Shared test doubles. Nothing here touches the network."""
from __future__ import annotations

import json
from typing import Any, AsyncIterator

from app.llm.client import Completed, Completion, Generated, HistoryMessage, LLMEvent, TextDelta, Usage


class ScriptedLLM:
    """A fake provider that plays back a fixed script of events / exceptions.

    Items in `script`: LLMEvent instances are yielded; Exception instances are
    raised at that point. That makes "stream 'Hello ', 'world', then crash"
    a one-liner, deterministic, and free.
    """

    model = "test-model"

    def __init__(self, script: list[LLMEvent | Exception], title: str | None = "Scripted Title",
                 model: str | None = None) -> None:
        if model:
            self.model = model  # instance attribute shadows the class default
        self.script = script
        self.title = title
        self.histories: list[list[HistoryMessage]] = []  # what the app sent us, per call
        self.upstream_closed = False  # did the app close the provider stream?
        self.title_requests: list[str] = []
        self.tools_seen: list = []  # the ToolExecutor (or None) passed for each reply
        self.cache_keys: list = []  # the prompt_cache_key passed for each reply
        self.summary_requests: list[tuple[str | None, list[HistoryMessage]]] = []  # (previous summary, folded messages)
        self.summary_text: str | None = "SCRIPTED SUMMARY"  # None = the summary call fails

    async def stream(self, history: list[HistoryMessage], tools=None, cache_key=None) -> AsyncIterator[LLMEvent]:
        self.histories.append(history)
        self.tools_seen.append(tools)
        self.cache_keys.append(cache_key)
        try:
            for item in self.script:
                if isinstance(item, Exception):
                    raise item
                yield item
        finally:
            self.upstream_closed = True

    async def complete(self, history: list[HistoryMessage]) -> Completion:
        self.histories.append(history)
        return Completion("complete-text", "completed", Usage(10, 5, 15))

    async def generate_title(self, first_user_text: str) -> Generated | None:
        self.title_requests.append(first_user_text)
        return Generated(self.title, Usage(30, 6, 36), "title-snapshot", self.model) if self.title else None

    async def summarize(self, previous_summary, messages) -> Generated | None:
        self.summary_requests.append((previous_summary, messages))
        if self.summary_text is None:
            return None
        return Generated(self.summary_text, Usage(200, 40, 240), "summary-snapshot", self.model)


def ok_script(*chunks: str, usage: Usage | None = Usage(100, 50, 150), stop: str = "completed",
              served_model: str | None = None):
    return [*(TextDelta(c) for c in chunks), Completed(stop, usage, served_model)]


def parse_sse(text: str) -> list[tuple[str, dict[str, Any]]]:
    """Parse a full SSE body into [(event, data_dict), ...] (test-side mirror of sse.ts)."""
    events = []
    for block in text.split("\n\n"):
        if not block.strip():
            continue
        name, data = "message", None
        for line in block.split("\n"):
            if line.startswith("event: "):
                name = line[len("event: "):]
            elif line.startswith("data: "):
                data = json.loads(line[len("data: "):])
        events.append((name, data))
    return events


# ── RAG test doubles ─────────────────────────────────────────────────────
import hashlib  # noqa: E402
import io  # noqa: E402
import math  # noqa: E402
import re  # noqa: E402

from app.rag.embeddings import EMBEDDING_DIMENSIONS, Embedded  # noqa: E402


def fake_vector(text: str) -> list[float]:
    """A deterministic stand-in for an embedding: each word is hashed to one of 1536 slots and the
    vector is normalised. Texts that SHARE WORDS get a high cosine similarity, which is enough to
    test retrieval end to end without a network call. (It has no real "meaning", unlike a model.)"""
    vec = [0.0] * EMBEDDING_DIMENSIONS
    for word in re.findall(r"[a-z0-9]+", text.lower()):
        vec[int(hashlib.md5(word.encode()).hexdigest(), 16) % EMBEDDING_DIMENSIONS] += 1.0
    norm = math.sqrt(sum(x * x for x in vec)) or 1.0
    return [x / norm for x in vec]


class FakeEmbedder:
    model = "fake-embedding"

    def __init__(self, fail_with: Exception | None = None) -> None:
        self.calls: list[list[str]] = []
        self.fail_with = fail_with

    async def embed(self, texts: list[str]) -> Embedded:
        self.calls.append(list(texts))
        if self.fail_with is not None:
            raise self.fail_with
        return Embedded([fake_vector(t) for t in texts], sum(len(t.split()) for t in texts), "fake-snapshot", self.model)


def make_pdf(pages: list[str]) -> bytes:
    """A tiny but VALID PDF with the given text, one string per page (so pypdf reads it like any PDF)."""
    objects: list[bytes] = []
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(len(pages)))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    font_id = 3 + 2 * len(pages)
    for i, text in enumerate(pages):
        content_id = 4 + 2 * i
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode()
        )
        safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 50 700 Td ({safe}) Tj ET".encode()
        objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def make_docx(blocks: list[tuple[str, str]], table: list[list[str]] | None = None) -> bytes:
    """A real .docx. `blocks` are (kind, text): kind is 'h1', 'h2' or 'p'."""
    from docx import Document

    document = Document()
    for kind, text in blocks:
        if kind.startswith("h"):
            document.add_heading(text, level=int(kind[1:]))
        else:
            document.add_paragraph(text)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, cell in enumerate(row):
                t.cell(r, c).text = cell
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
